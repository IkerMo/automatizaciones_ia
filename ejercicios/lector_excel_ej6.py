from openpyxl import Workbook, load_workbook

wb = load_workbook(filename='inputs/sys_ejemplo.xlsx', data_only=True)
ws = wb.active

for row in ws.iter_rows():
    for cell in row:
        if (cell.value.lower == 'código'): break
